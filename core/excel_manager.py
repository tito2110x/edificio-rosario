import openpyxl
from openpyxl.utils import get_column_letter
import os
import shutil
from datetime import datetime
from config import DEFAULT_EXCEL_PATH, DEPARTAMENTOS

def obtener_ultimas_lecturas(excel_path: str = DEFAULT_EXCEL_PATH) -> dict:
    """
    Lee el archivo Excel y extrae las lecturas más recientes para cada dpto.
    Retorna: {
        'nombre_mes': 'Septiembre',
        'lecturas': {'101': 155691.0, ...},
        'soles': {'101': 75.41, ...}
    }
    """
    if not os.path.exists(excel_path):
        return {"nombre_mes": "", "lecturas": {}, "soles": {}}

    wb = openpyxl.load_workbook(excel_path, data_only=True)
    if "AGUA" not in wb.sheetnames:
        return {"nombre_mes": "", "lecturas": {}, "soles": {}}

    ws = wb["AGUA"]
    max_col = ws.max_column

    # Buscar la última columna de 'Lectura actual'
    ultima_col_lectura = None
    ultima_col_soles = None
    nombre_mes = ""

    for c in range(max_col, 1, -1):
        header_val = str(ws.cell(4, c).value or "").lower()
        if "lectura" in header_val and ultima_col_lectura is None:
            ultima_col_lectura = c
            # Intentar ver el nombre del mes en fila 3 o 1
            mes_r3 = ws.cell(3, c).value
            mes_r1 = ws.cell(1, c).value
            nombre_mes = str(mes_r3 or mes_r1 or f"Columna {c}")

        if "soles" in header_val and ultima_col_soles is None and ultima_col_lectura is not None:
            # Si ya encontramos la lectura, la columna de soles suele estar 2 columnas a la derecha
            pass

    # Si encontramos ultima_col_lectura
    if ultima_col_lectura is None:
        ultima_col_lectura = 13  # Fallback a Columna M

    # Buscar la columna de soles de ese mismo mes
    for c in range(ultima_col_lectura, min(ultima_col_lectura + 4, max_col + 1)):
        h = str(ws.cell(4, c).value or "").lower()
        if "soles" in h:
            ultima_col_soles = c
            break

    lecturas = {}
    soles = {}

    # Mapear filas por departamento (Filas 5 a 16)
    for r in range(5, 17):
        dpto_cell = str(ws.cell(r, 1).value or "")
        # Extraer número de dpto (101, 102...)
        for d in DEPARTAMENTOS:
            if d["dpto"] in dpto_cell:
                val_lectura = ws.cell(r, ultima_col_lectura).value
                val_soles = ws.cell(r, ultima_col_soles).value if ultima_col_soles else 0.0
                try:
                    lecturas[d["dpto"]] = float(val_lectura or 0.0)
                except (ValueError, TypeError):
                    lecturas[d["dpto"]] = 0.0
                try:
                    soles[d["dpto"]] = float(val_soles or 0.0)
                except (ValueError, TypeError):
                    soles[d["dpto"]] = 0.0
                break

    return {
        "nombre_mes": nombre_mes,
        "col_lectura": ultima_col_lectura,
        "col_soles": ultima_col_soles,
        "lecturas": lecturas,
        "soles": soles
    }

def agregar_nuevo_mes_excel(
    nombre_nuevo_mes: str,
    monto_recibo_sedapal: float,
    lecturas_actuales: dict[str, float],
    excel_path: str = DEFAULT_EXCEL_PATH,
    output_path: str = None
) -> str:
    """
    Inserta las nuevas 4 columnas en la hoja AGUA y actualiza RESUMEN
    respetando exactamente la estructura y fórmulas de Excel.
    """
    if not os.path.exists(excel_path):
        raise FileNotFoundError(f"No se encontró el archivo Excel en {excel_path}")

    if output_path is None:
        # Guardar en el mismo archivo creando respaldo previo
        backup_path = excel_path.replace(".xlsx", f"_backup_{datetime.now().strftime('%Y%m%d_%H%M%S')}.xlsx")
        shutil.copyfile(excel_path, backup_path)
        output_path = excel_path

    # Cargar workbook con fórmulas
    wb = openpyxl.load_workbook(excel_path, data_only=False)
    ws_agua = wb["AGUA"]

    info_anterior = obtener_ultimas_lecturas(excel_path)
    col_ant_lectura = info_anterior["col_lectura"]
    col_ant_soles = info_anterior["col_soles"]

    # Determinar nueva columna inicial
    max_col = ws_agua.max_column
    col_nueva_lectura = max_col + 1
    col_consumo = col_nueva_lectura + 1
    col_soles = col_consumo + 1
    col_variacion = col_soles + 1

    letra_ant_lectura = get_column_letter(col_ant_lectura)
    letra_ant_soles = get_column_letter(col_ant_soles) if col_ant_soles else "L"

    letra_nueva_lectura = get_column_letter(col_nueva_lectura)
    letra_consumo = get_column_letter(col_consumo)
    letra_soles = get_column_letter(col_soles)
    letra_variacion = get_column_letter(col_variacion)

    # Encabezados
    ws_agua.cell(1, col_nueva_lectura, f"Total Recibo {nombre_nuevo_mes}")
    ws_agua.cell(2, col_nueva_lectura, float(monto_recibo_sedapal))
    ws_agua.cell(3, col_nueva_lectura, nombre_nuevo_mes)

    ws_agua.cell(4, col_nueva_lectura, "Lectura actual")
    ws_agua.cell(4, col_consumo, "Consumo m3")
    ws_agua.cell(4, col_soles, f"Total soles {nombre_nuevo_mes}")
    ws_agua.cell(4, col_variacion, "Variación mes anterior %")

    # Filas 5 a 16 (Departamentos)
    for r in range(5, 17):
        dpto_cell = str(ws_agua.cell(r, 1).value or "")
        dpto_num = None
        for d in DEPARTAMENTOS:
            if d["dpto"] in dpto_cell:
                dpto_num = d["dpto"]
                break

        val_act = float(lecturas_actuales.get(dpto_num, 0.0)) if dpto_num else 0.0

        # Escribir lectura
        ws_agua.cell(r, col_nueva_lectura, val_act)
        # Fórmula consumo: =+NuevaLectura - AntLectura
        ws_agua.cell(r, col_consumo, f"=+{letra_nueva_lectura}{r}-{letra_ant_lectura}{r}")
        # Fórmula soles: =+([Consumo]/$[TotalConsumo])*$[TotalRecibo]
        ws_agua.cell(r, col_soles, f"=({letra_consumo}{r}/${letra_consumo}$17)*${letra_nueva_lectura}$2")
        # Fórmula variación: =+[Soles]/[SolesAnterior]-1
        ws_agua.cell(r, col_variacion, f"=+{letra_soles}{r}/{letra_ant_soles}{r}-1")

    # Fila 17 (Totales)
    ws_agua.cell(17, col_consumo, f"=+SUM({letra_consumo}5:{letra_consumo}16)")
    ws_agua.cell(17, col_soles, f"=+SUM({letra_soles}5:{letra_soles}16)")

    # Actualizar hoja RESUMEN para apuntar a la nueva columna de soles
    if "RESUMEN" in wb.sheetnames:
        ws_resumen = wb["RESUMEN"]
        for r in range(6, 18):
            row_idx = r - 1  # fila 6 de RESUMEN corresponde a fila 5 de AGUA
            ws_resumen.cell(r, 2, f"=+AGUA!{letra_soles}{row_idx}")

    wb.save(output_path)
    return output_path
