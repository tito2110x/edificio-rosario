from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
import io
from config import DEPARTAMENTOS

def generar_powerpoint_mensual(
    nombre_mes: str,
    resumen_calculos: list[dict],
    fotos_por_dpto: dict[str, bytes],
    suministro_sedapal: str = "2675620"
) -> bytes:
    """
    Genera un archivo PowerPoint (.pptx) idéntico en estructura al de Rosario del Solar,
    con la lámina de portada y una lámina por departamento con su foto y datos.
    """
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # --- Lámina 1: Portada ---
    slide1 = prs.slides.add_slide(blank_layout)
    txBox = slide1.shapes.add_textbox(Inches(1.5), Inches(2.5), Inches(10.33), Inches(2.5))
    tf = txBox.text_frame
    p = tf.paragraphs[0]
    p.text = f"CONSUMO DE AGUA - {nombre_mes.upper()}"
    p.font.size = Pt(44)
    p.font.bold = True
    p.font.color.rgb = RGBColor(14, 74, 133)

    p2 = tf.add_paragraph()
    p2.text = f"Edificio Rosario del Solar | Suministro Agua: {suministro_sedapal}"
    p2.font.size = Pt(24)
    p2.font.color.rgb = RGBColor(80, 80, 80)

    # --- Láminas 2 a 13: Por Departamento ---
    mapa_datos = {d["dpto"]: d for d in resumen_calculos}

    for dpto_info in DEPARTAMENTOS:
        dpto = dpto_info["dpto"]
        nombre = dpto_info["nombre"]
        data = mapa_datos.get(dpto, {})

        slide = prs.slides.add_slide(blank_layout)

        # Título
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(1.0))
        tf = tb.text_frame
        p = tf.paragraphs[0]
        p.text = f"CONSUMO DE AGUA {dpto} - {nombre.upper()}"
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.color.rgb = RGBColor(14, 74, 133)

        # Cuadro de datos
        tb_data = slide.shapes.add_textbox(Inches(0.8), Inches(1.8), Inches(5.5), Inches(4.5))
        tf_data = tb_data.text_frame
        
        items = [
            ("Lectura Anterior", f"{data.get('lectura_anterior', 0):,.0f}"),
            ("Lectura Actual", f"{data.get('lectura_actual', 0):,.0f}"),
            ("Consumo del Mes", f"{data.get('consumo', 0):,.0f} m³"),
            ("Importe Agua", f"S/ {data.get('agua_soles', 0):.2f}"),
            ("Cuota Mantenimiento", f"S/ {data.get('cuota_mantenimiento', 30):.2f}"),
            ("Total a Pagar", f"S/ {data.get('total_a_pagar', 0):.2f}"),
        ]

        for i, (label, val) in enumerate(items):
            p = tf_data.add_paragraph() if i > 0 else tf_data.paragraphs[0]
            p.text = f"• {label}: {val}"
            p.font.size = Pt(20)
            if "Total" in label:
                p.font.bold = True
                p.font.color.rgb = RGBColor(180, 30, 30)
                p.font.size = Pt(22)
            else:
                p.font.color.rgb = RGBColor(40, 40, 40)

        # Imagen del medidor
        img_bytes = fotos_por_dpto.get(dpto)
        if img_bytes:
            try:
                img_stream = io.BytesIO(img_bytes)
                slide.shapes.add_picture(img_stream, Inches(6.8), Inches(1.6), width=Inches(5.5))
            except Exception:
                pass

    out = io.BytesIO()
    prs.save(out)
    return out.getvalue()
